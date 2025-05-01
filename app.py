import streamlit as st
import pandas as pd
import re
from openpyxl import load_workbook
from docx import Document
import tempfile
import os

st.set_page_config(page_title="Conversor para Excel e Word", layout="centered")

st.title("Conversor Automático para .xlsx e .docx")
st.write("Este app formata automaticamente sua planilha e gera uma versão para Word com a mesma tabela.")

uploaded_file = st.file_uploader("Envie sua planilha (.xlsx ou .csv)", type=["xlsx", "csv"])

def formatar_texto(texto):
    if pd.isnull(texto):
        return texto
    texto = str(texto).strip().strip('"').strip("'")
    texto = re.sub(r'\s+', ' ', texto)
    frases = re.split(r'(?<=[.!?]) +', texto)
    frases = [f[:1].upper() + f[1:].lower() + ('' if f.strip().endswith('.') else '.') for f in frases if f]
    return ' '.join(frases)

if uploaded_file:
    ext = os.path.splitext(uploaded_file.name)[-1].lower()
    if ext == '.xlsx':
        df = pd.read_excel(uploaded_file)
    elif ext == '.csv':
        df = pd.read_csv(uploaded_file)
    else:
        st.error("Formato não suportado. Use .xlsx ou .csv")

    coluna_c_nome = df.columns[2]
    df[coluna_c_nome] = df[coluna_c_nome].apply(formatar_texto)

    # Corrigir cabeçalhos
    df.columns = [col.upper().replace("QTDE", "QUANTIDADE") if "CATMAT" not in col.upper() else col for col in df.columns]

    # Salvar Excel
    with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as tmp_excel:
        df.to_excel(tmp_excel.name, index=False)
        wb = load_workbook(tmp_excel.name)
        ws = wb.active
        for row in range(2, ws.max_row + 1):
            for col in ['F', 'G']:
                try:
                    cell = ws[f"{col}{row}"]
                    cell.value = float(cell.value)
                    cell.number_format = '"R$"#,##0.00'
                except:
                    pass
        wb.save(tmp_excel.name)
        st.download_button("Baixar Excel Formatado", tmp_excel.read(), file_name="planilha_formatada.xlsx")

    # Salvar Word
    with tempfile.NamedTemporaryFile(delete=False, suffix=".docx") as tmp_doc:
        doc = Document()
        doc.add_heading("Tabela da Planilha", level=1)
        tabela = doc.add_table(rows=1, cols=len(df.columns))
        tabela.style = 'Table Grid'
        hdr_cells = tabela.rows[0].cells
        for i, coluna in enumerate(df.columns):
            hdr_cells[i].text = str(coluna)
        for _, row in df.iterrows():
            linha = tabela.add_row().cells
            for i, valor in enumerate(row):
                linha[i].text = str(valor)
        doc.save(tmp_doc.name)
        st.download_button("Baixar Word com Tabela", tmp_doc.read(), file_name="tabela_formatada.docx")